import csv
import os
import tempfile
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, datetime, time as dt_time, timedelta
from pathlib import Path
from typing import Dict, List, Optional, Set

import pythoncom
from openpyxl import load_workbook
from win32com.client import DispatchEx

from merge_files import merge_files_in_folder

BASE_DIR = Path(__file__).resolve().parent
INPUT_FOLDER = BASE_DIR / "Daroda_lc_vrn_files/vrn"
MERGE_OUTPUT_FILE = BASE_DIR / "normalized_and_merged_vrn_daroda.csv"

MERGE_AFTER_NORMALIZATION = True
# Each .xlsx is fully loaded by openpyxl; parallel workers multiply peak RAM (workers × workbook size).
MAX_WORKER_THREADS = 5
# Before normalization: try csv_converter.convert_workbook on each .xlsx/.xls; on success
# delete the workbook and keep the new .csv. Failures keep the original for default handling.

MERGE_HEADER_KEYWORDS = [
    "MVC_TLC_CLASS",
    "CCH TXN NO",
    "AVC",
    "TRANSACTION NO",
    "MVC",
    "VEH CLASS",
    "TC CLASS",
    "Veh Class",
    "OperatorClass",
    "TcClass",
    "Operator Class",
    "MVC TLC CLASS",
]

SUPPORTED_EXTENSIONS = {".xlsx", ".xls", ".csv"}
TRACKED_HEADER_COLUMNS = {
    "TC Class",
    "Veh Reg No.",
    "MOP",
    "Lane No",
    "Description",
    "File Name",
    "Date & Time",
}

NORMALIZATION_GROUPS = {
    "TC Class": [
        "MVC_TLC_CLASS",
        "MVC",
        "MVC (TLC CLASS)",
        "VEH CLASS",
        "TC CLASS",
        "Veh Class",
        "Operator_Class",
        "OperatorClass",
        "TcClass",
        "Operator Class",
        "MVC TLC CLASS",
    ],
    "Veh Reg No.": [
        "TC_VEH_REG_NO",
        "Vehicle Reg. No.",
        "Vehicle Reg. No",
        "VEHICLE_REG_NO",
        "veh_reg_no_",
        "Veh Reg No.",
        "VEH REG NO",
        "TC VEH REG NO",
        "VEH. REG. NO.",
        "Licence Plate No",
        "VRN",
        "Licence Plate No.",
        "Plate No",
        "NPCI VRN",
        "Veh Reg Num",
        "VehicleNumber",
        "VEH_REG_NO",
        "Platenumber",
        "Vehicle Registration Number",
        "vehicle_reg_no",
        "Vehicle No",
    ],
    "MOP": [
        "PAYMENT_TYPE",
        "Payment Method",
        "MVC MOP",
        "PAYMENT TYPE",
        "Payment",
        "Mode",
        "Ticket_Type",
        "MVC_TLC_MOP",
        "TransactionTypeTC",
        "PaymentMeans",
        "PAYMENT METHOD",
        "MVC (TLC MOP)",
        "MVC TLC MOP",
    ],
    "Lane No": [
        "LANE NO",
        "LaneNo",
        "Lane ID",
        "Lane Number",
        "LANE_NUMBER",
        "Lane",
    ],
    "Description": [
        "DESC",
        "DESCRIPTION",
        "Remarks",
        "Remark"
    ],
    "File Name": [
        "FILENAME",
        "FileName",
        "SOURCE FILE",
    ],
    "Date & Time": [
        "DATE",
        "Date Time",
        "Reader Read Time",
        "DATETIME",
        "DATE TIME",
        "Date/Time",
        "Transaction Date",
        "Txn Date",
        "TXN DATE",
    ],
    "Car": ["CarJeep", "CAR/JEEP/VAN", r"CAR\JEEP"],
    "LCV": ["MiniBus", "LCV/MINI BUS"],
    "MAV": [
        "MAV_4",
        "TRUCK 4-6 AXLE",
    ],
    "TRUCK 3 AXLE": [
        "TRUCK 3 AXLE-2T",
        "Truck3X",
        "TRUCK-3 AXLE",
    ],
    "TRK 2 AXLE": [
        "TRUCK-2 AXLE",
        "BUS-2 AXLE",
        "TRUCK 2 AXLE",
    ],
    "TAG": [
        "TAG-",
        "TAG-NA",
        "TAGNA",
    ],
    "CASH": [
        "CASH-UPI",
        "CASHS",
    ],
    "ETC": [
        "ETC_PENALTY_CASH",
        "ETC_PENALTY_CD",
    ],
    "EXEMPT": [
        "EXEMPTED",
    ],
}


def normalize_text(value):
    if value is None:
        return ""
    text = str(value).strip().casefold()
    for char in ("_", "-", "/", "\\", ".", ",", "(", ")", "[", "]"):
        text = text.replace(char, " ")
    return " ".join(text.split())


def build_lookup(groups):
    lookup = {}
    for common_name, aliases in groups.items():
        lookup[normalize_text(common_name)] = common_name
        for alias in aliases:
            lookup[normalize_text(alias)] = common_name
    return lookup


def update_value(value, lookup):
    replacement = lookup.get(normalize_text(value))
    if replacement and value != replacement:
        return replacement
    return None


def normalize_vrn_value(value):
    if value is None:
        return None

    text = str(value).strip().upper()
    if not text:
        return None

    normalized = "".join(ch for ch in text if ch.isalnum())
    if not normalized or normalized == text:
        return None

    return normalized


def clean_lane_value(value):
    if value is None:
        return None

    text = str(value)
    if not text:
        return None

    transformed = text
    transformed = transformed.replace("Lane ", "L0")
    transformed = transformed.replace("L010", "L10")
    transformed = transformed.replace("LN00", "L")
    transformed = transformed.replace("LN0", "L")
    transformed = transformed.replace("LN", "L")
    transformed = transformed.replace("L011", "L11")
    transformed = transformed.replace("L012", "L12")
    transformed = transformed.split("-", 1)[0]
    transformed = transformed.strip()
    transformed = transformed.replace("L0", "L")

    if transformed != text:
        return transformed
    return None


def clean_description_value(value):
    if value is None:
        return None

    text = str(value)
    if not text:
        return None

    transformed = text
    transformed = transformed.replace("()", "")
    transformed = transformed.replace("(CASH)", "")
    transformed = transformed.replace(" - DOWN", "")
    transformed = transformed.replace(" - UP", "")
    transformed = transformed.replace("ED - ", "")
    transformed = transformed.replace("EXEMPT", "")

    if transformed != text:
        return transformed
    return None


def clean_mop_value(value):
    if value is None:
        return None

    text = str(value)
    if not text:
        return None

    transformed = text.upper().split("-", 1)[0]
    if transformed != text:
        return transformed
    return None


def clean_tc_class_value(value):
    if value is None:
        return None

    text = str(value)
    if not text:
        return None

    if "MAV" in text.upper() and text != "MAV":
        return "MAV"
    return None


def _excel_serial_to_datetime(serial: float) -> datetime:
    base = datetime(1899, 12, 30)
    return base + timedelta(days=float(serial))


def parse_datetime_cell(value) -> Optional[datetime]:
    """Parse a cell value to datetime (Power BI datetime / text / Excel serial)."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date) and not isinstance(value, datetime):
        return datetime.combine(value, dt_time.min)
    if isinstance(value, (int, float)):
        try:
            return _excel_serial_to_datetime(float(value))
        except (ValueError, OSError, OverflowError):
            return None
    text = str(value).strip()
    if not text:
        return None
    fmts = (
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%d-%m-%Y %H:%M:%S",
        "%d-%m-%Y %H:%M",
        "%d/%m/%Y %H:%M:%S",
        "%d/%m/%Y %H:%M",
        "%m/%d/%Y %H:%M:%S",
        "%m/%d/%Y %H:%M",
        "%Y-%m-%d",
        "%d-%m-%Y",
        "%d/%m/%Y",
        "%m/%d/%Y",
    )
    for fmt in fmts:
        try:
            return datetime.strptime(text[:26], fmt)
        except ValueError:
            continue
    return None


def coerce_datetime_value(value) -> Optional[datetime]:
    """Power Query Changed Type to datetime: normalize parseable non-datetime cells."""
    if isinstance(value, datetime):
        return None
    return parse_datetime_cell(value)


def resolve_header_name(value, normalization_lookup):
    canonical = normalization_lookup.get(normalize_text(value))
    if canonical in TRACKED_HEADER_COLUMNS:
        return canonical
    return None


def should_drop_row(row_values, header_columns):
    lane_col_idx = next((idx for idx, name in header_columns.items() if name == "Lane No"), None)
    if lane_col_idx is None or lane_col_idx > len(row_values):
        return False

    lane_value = row_values[lane_col_idx - 1]
    if lane_value is None or str(lane_value).strip() == "":
        return True
    if str(lane_value).strip().upper() == "EXEMPTED":
        return True
    return False


def process_row_values(row_values, normalization_lookup, header_columns: Dict[int, str]):
    updates = []
    row_values_out = list(row_values)
    header_cell_indexes: Set[int] = set()

    for col_idx, value in enumerate(row_values, start=1):
        header_name = resolve_header_name(value, normalization_lookup)
        if header_name is not None:
            header_columns[col_idx] = header_name
            header_cell_indexes.add(col_idx)

        replacement = update_value(value, normalization_lookup)
        if replacement is not None:
            updates.append((col_idx, replacement))
            row_values_out[col_idx - 1] = replacement

    for col_idx, value in enumerate(row_values_out, start=1):
        if col_idx in header_cell_indexes:
            continue

        column_name = header_columns.get(col_idx)
        replacement = None

        if column_name == "Date & Time":
            replacement = coerce_datetime_value(value)
        elif column_name == "Veh Reg No.":
            replacement = normalize_vrn_value(value)
        elif column_name == "MOP":
            replacement = clean_mop_value(value)
        elif column_name == "Lane No":
            replacement = clean_lane_value(value)
        elif column_name == "Description":
            replacement = clean_description_value(value)
        elif column_name == "TC Class":
            replacement = clean_tc_class_value(value)

        if replacement is not None and replacement != value:
            updates.append((col_idx, replacement))
            row_values_out[col_idx - 1] = replacement

    drop_row = should_drop_row(row_values_out, header_columns)
    return updates, drop_row


def get_last_used_row_col(ws):
    xlFormulas = -4123
    xlByRows = 1
    xlByColumns = 2
    xlPrevious = 2

    last_row_cell = ws.Cells.Find(
        What="*", LookIn=xlFormulas, SearchOrder=xlByRows, SearchDirection=xlPrevious
    )

    last_col_cell = ws.Cells.Find(
        What="*", LookIn=xlFormulas, SearchOrder=xlByColumns, SearchDirection=xlPrevious
    )

    if last_row_cell is None or last_col_cell is None:
        return 0, 0

    return last_row_cell.Row, last_col_cell.Column


def normalize_xlsx_file(file_path: Path, normalization_lookup) -> int:
    print(f"[START][XLSX] {file_path.name}")
    workbook = load_workbook(file_path, keep_links=False)
    total_changes = 0

    try:
        for ws in workbook.worksheets:
            print(
                f"[SCAN] {file_path.name} -> {ws.title}: single-pass worksheet traversal"
            )
            sheet_changes = 0
            header_columns: Dict[int, str] = {}
            rows_to_delete: List[int] = []

            for row_idx, row in enumerate(
                ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column),
                start=1,
            ):
                row_values = [cell.value for cell in row]
                updates, drop_row = process_row_values(
                    row_values,
                    normalization_lookup,
                    header_columns,
                )

                if drop_row:
                    rows_to_delete.append(row_idx)
                    continue

                for col_idx, replacement in updates:
                    row[col_idx - 1].value = replacement
                    total_changes += 1
                    sheet_changes += 1

            for row_idx in reversed(rows_to_delete):
                ws.delete_rows(row_idx, 1)
                total_changes += 1
                sheet_changes += 1

            file_name_col = next(
                (idx for idx, name in header_columns.items() if name == "File Name"),
                None,
            )
            if file_name_col is not None:
                ws.delete_cols(file_name_col, 1)
                total_changes += 1
                sheet_changes += 1

            if sheet_changes:
                print(
                    f"[MATCH] {file_path.name} -> {ws.title}: "
                    f"{sheet_changes} cell(s) normalized"
                )
            else:
                print(f"[NO-MATCH] {file_path.name} -> {ws.title}")

        if total_changes:
            workbook.save(file_path)
            print(f"[SAVE][XLSX] {file_path.name}")
        else:
            print(f"[SKIP][XLSX] No changes: {file_path.name}")
    finally:
        workbook.close()

    return total_changes


def normalize_xls_file(file_path: Path, normalization_lookup) -> int:
    pythoncom.CoInitialize()

    excel = None
    workbook = None

    try:
        print(f"[START][XLS] {file_path.name}")

        excel = DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        try:
            excel.ScreenUpdating = False
        except Exception:
            pass

        try:
            excel.EnableEvents = False
        except Exception:
            pass

        try:
            excel.Calculation = -4135
        except Exception:
            pass  # manual calculation

        workbook = excel.Workbooks.Open(str(file_path))

        total_changes = 0

        for ws in workbook.Worksheets:
            last_row, last_col = get_last_used_row_col(ws)

            if last_row == 0 or last_col == 0:
                print(f"[SKIP-SHEET] {file_path.name} -> {ws.Name}: empty sheet")
                continue

            print(
                f"[SCAN] {file_path.name} -> {ws.Name}: "
                f"single-pass rows 1-{last_row}, columns 1-{last_col}"
            )

            sheet_changes = 0
            header_columns: Dict[int, str] = {}
            rows_to_delete: List[int] = []

            for r_idx in range(1, last_row + 1):
                row_values = [
                    ws.Cells(r_idx, c_idx).Text for c_idx in range(1, last_col + 1)
                ]
                updates, drop_row = process_row_values(
                    row_values,
                    normalization_lookup,
                    header_columns,
                )

                if drop_row:
                    rows_to_delete.append(r_idx)
                    continue

                for c_idx, replacement in updates:
                    cell = ws.Cells(r_idx, c_idx)
                    cell.Value = replacement
                    total_changes += 1
                    sheet_changes += 1

            for r_idx in reversed(rows_to_delete):
                ws.Rows(r_idx).Delete()
                total_changes += 1
                sheet_changes += 1

            file_name_col = next(
                (idx for idx, name in header_columns.items() if name == "File Name"),
                None,
            )
            if file_name_col is not None:
                ws.Columns(file_name_col).Delete()
                total_changes += 1
                sheet_changes += 1

            if sheet_changes:
                print(
                    f"[MATCH] {file_path.name} -> {ws.Name}: "
                    f"{sheet_changes} cell(s) normalized"
                )
            else:
                print(f"[NO-MATCH] {file_path.name} -> {ws.Name}")

        if total_changes:
            workbook.Save()
            print(f"[SAVE] {file_path.name}")
        else:
            print(f"[SKIP][XLS] No changes: {file_path.name}")

        workbook.Close(SaveChanges=False)
        excel.Quit()

        return total_changes

    finally:
        try:
            if workbook is not None:
                workbook.Close(SaveChanges=False)
        except Exception:
            pass

        try:
            if excel is not None:
                excel.Quit()
        except Exception:
            pass

        pythoncom.CoUninitialize()


def detect_csv_encoding(file_path: Path):
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            with file_path.open("r", encoding=enc, newline="") as f:
                f.read(4096)
            return enc
        except UnicodeDecodeError:
            continue
    return "latin-1"


def normalize_csv_file_fast(file_path: Path, normalization_lookup):
    print(
        f"[START][CSV] {file_path.name}: detect encoding, stream normalize, write temp",
        flush=True,
    )
    encoding = detect_csv_encoding(file_path)
    print(
        f"[OPER][CSV] {file_path.name}: encoding={encoding}; "
        "normalize cells / drop rows / strip File Name column",
        flush=True,
    )
    total_changes = 0

    temp_fd, temp_path = tempfile.mkstemp(suffix=".csv")
    os.close(temp_fd)

    with (
        file_path.open("r", encoding=encoding, newline="") as src,
        open(temp_path, "w", encoding=encoding, newline="") as dst,
    ):
        reader = csv.reader(src)
        writer = csv.writer(dst)
        header_columns: Dict[int, str] = {}
        file_name_col_idx = None

        for row in reader:
            updates, drop_row = process_row_values(
                row,
                normalization_lookup,
                header_columns,
            )

            for col_index, replacement in updates:
                row[col_index - 1] = replacement
                total_changes += 1

            if drop_row:
                total_changes += 1
                continue

            if file_name_col_idx is None:
                file_name_col_idx = next(
                    (idx for idx, name in header_columns.items() if name == "File Name"),
                    None,
                )
            if file_name_col_idx is not None and file_name_col_idx <= len(row):
                del row[file_name_col_idx - 1]
                total_changes += 1

            writer.writerow(row)

    if total_changes:
        os.replace(temp_path, file_path)
        print(
            f"[SAVE][CSV] {file_path.name}: replace file in place (edits applied)",
            flush=True,
        )
    else:
        os.remove(temp_path)
        print(
            f"[SKIP][CSV] {file_path.name}: discard temp (no edits)",
            flush=True,
        )

    return total_changes


def find_input_files(folder_path: Path):
    return sorted(
        p
        for p in folder_path.iterdir()
        if p.is_file() and p.suffix.lower() in SUPPORTED_EXTENSIONS
    )


def phase_convert_workbooks_to_csv_and_delete(folder_path: Path) -> None:
    """
    Use csv_converter on each .xlsx/.xls in folder. On success, remove the workbook
    so phase 2 only normalizes CSV where conversion worked; failures keep originals.
    """
    try:
        from csv_converter import (
            HEADER_KEYWORDS as _csv_header_keywords,
            HEADER_SCAN_MAX_ROW as _csv_scan_max_row,
            MIN_HEADER_KEYWORD_MATCHES as _csv_min_matches,
            convert_workbook,
        )
    except ImportError as exc:
        print(
            f"[WARN] Could not import csv_converter ({exc}); "
            "skipping pre-conversion to CSV.",
            flush=True,
        )
        return

    keywords = list(_csv_header_keywords)
    if len(keywords) < _csv_min_matches:
        print(
            f"[WARN] csv_converter needs at least {_csv_min_matches} HEADER_KEYWORDS; "
            f"have {len(keywords)}. Skipping pre-conversion to CSV.",
            flush=True,
        )
        return

    folder_path = folder_path.resolve()
    if not folder_path.is_dir():
        return

    targets = [
        p
        for p in sorted(folder_path.iterdir(), key=lambda x: x.name.lower())
        if p.is_file()
        and p.suffix.lower() in (".xlsx", ".xls")
        and not p.name.startswith("~$")
    ]
    if not targets:
        print("[PHASE 1] No .xlsx/.xls files to pre-convert.", flush=True)
        return

    print(
        f"\n[PHASE 1] csv_converter: trying {len(targets)} workbook(s) → .csv "
        f"(header scan rows 1-{_csv_scan_max_row}, min matches={_csv_min_matches}; "
        f"delete source on success)…",
        flush=True,
    )
    for path in targets:
        out_csv = path.with_suffix(".csv")
        try:
            print(
                f"[PHASE 1][START] {path.name}: detecting header + converting...",
                flush=True,
            )
            convert_workbook(path, keywords, None)
            path.unlink()
            print(
                f"[PHASE 1][DONE] {path.name}: converted to {out_csv.name}",
                flush=True,
            )
            print(f"[CONVERT+DEL] {path.name} → {out_csv.name}", flush=True)
        except Exception as exc:
            print(
                f"[CONVERT-SKIP] {path.name}: {exc} "
                f"(will use default .xlsx/.xls normalization)",
                flush=True,
            )


def merge_normalized_files():
    print(f"\n[MERGE] Starting merge from {INPUT_FOLDER}")
    merge_files_in_folder(
        str(INPUT_FOLDER),
        str(MERGE_OUTPUT_FILE),
        MERGE_HEADER_KEYWORDS,
    )


def process_file(file_path: Path, normalization_lookup):
    suffix = file_path.suffix.lower()

    if suffix == ".xlsx":
        total = normalize_xlsx_file(file_path, normalization_lookup)
    elif suffix == ".xls":
        total = normalize_xls_file(file_path, normalization_lookup)
    elif suffix == ".csv":
        total = normalize_csv_file_fast(file_path, normalization_lookup)
    else:
        return None

    return file_path.name, total


def main():
    start_time = time.perf_counter()
    normalization_lookup = build_lookup(NORMALIZATION_GROUPS)

    phase_convert_workbooks_to_csv_and_delete(INPUT_FOLDER)

    files = find_input_files(INPUT_FOLDER)

    print(f"\n[PHASE 2] Found {len(files)} file(s) to normalize")

    worker_count = min(MAX_WORKER_THREADS, len(files)) if files else 0
    print(f"Using up to {worker_count} concurrent worker(s)")

    with ThreadPoolExecutor(max_workers=MAX_WORKER_THREADS) as executor:
        future_to_file = {
            executor.submit(process_file, file_path, normalization_lookup): file_path
            for file_path in files
        }

        for future in as_completed(future_to_file):
            file_path = future_to_file[future]

            try:
                result = future.result()
                if result is None:
                    continue

                file_name, total = result
                print(f"\nFile: {file_name}")
                print(f"Total cells changed: {total}")

            except Exception as e:
                print(f"\nError in {file_path.name}: {e}")

    if MERGE_AFTER_NORMALIZATION:
        merge_normalized_files()
    else:
        print("\n[MERGE] Skipped because MERGE_AFTER_NORMALIZATION is False")

    elapsed_seconds = time.perf_counter() - start_time
    print(f"[TIME] Total time taken: {elapsed_seconds:.2f} seconds")


if __name__ == "__main__":
    main()
